"""
generate_report_py.py
──────────────────────
生成 PHC-lite HLS 技術報告 Word 文件

安裝依賴（只需執行一次）：
    pip install python-docx

執行：
    python generate_report_py.py

輸出：PHC_lite_HLS_Report.docx（同目錄下）
"""

from __future__ import annotations
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ─── 顏色常數 ─────────────────────────────────────────────────────────────────
BLUE   = RGBColor(0x2E, 0x5F, 0xA3)
DGRAY  = RGBColor(0x40, 0x40, 0x40)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY  = RGBColor(0xF2, 0xF2, 0xF2)
GREEN  = RGBColor(0x2E, 0x7D, 0x32)

# ─── 建立文件 ─────────────────────────────────────────────────────────────────
doc = Document()

# 邊距 2.5 cm
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

def set_run_fmt(run, size_pt: float, bold=False, italic=False,
                color: RGBColor | None = None, font: str = "Arial"):
    run.font.name  = font
    run.font.size  = Pt(size_pt)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = BLUE
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(18)

def h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = DGRAY
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(14)

def h3(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=3)
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(12)

def para(doc: Document, text: str, bold=False, italic=False,
         color: RGBColor | None = None) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_fmt(run, 11, bold=bold, italic=italic, color=color)

def para_mixed(doc: Document, parts: list[tuple[str, dict]]) -> None:
    """parts = [(text, {bold, italic, color}), ...]"""
    p = doc.add_paragraph()
    for text, opts in parts:
        run = p.add_run(text)
        set_run_fmt(run, 11,
                    bold=opts.get("bold", False),
                    italic=opts.get("italic", False),
                    color=opts.get("color", None))

def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_fmt(run, 11)

def shade_cell(cell, hex_color: str) -> None:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def make_table(doc: Document, headers: list[str],
               rows: list[list[str]], col_widths_cm: list[float]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"

    # Header row
    hdr = t.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(col_widths_cm[i])
        shade_cell(cell, "2E5FA3")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_fmt(run, 10, bold=True, color=WHITE)

    # Data rows
    for ri, row_data in enumerate(rows):
        row = t.rows[ri + 1]
        fill = "FFFFFF" if ri % 2 == 0 else "F2F2F2"
        for ci, text in enumerate(row_data):
            cell = row.cells[ci]
            cell.width = Cm(col_widths_cm[ci])
            shade_cell(cell, fill)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            set_run_fmt(run, 10)

    doc.add_paragraph()  # 表格後空行


def page_break(doc: Document) -> None:
    doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# 封面
# ═══════════════════════════════════════════════════════════════════════════════
for _ in range(4):
    doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title_p.add_run("PHC-lite Neural Network Inference on FPGA")
set_run_fmt(r, 24, bold=True, color=BLUE)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sub_p.add_run("High-Frequency Trading FPGA 競賽專案  |  完整技術報告")
set_run_fmt(r2, 14, color=DGRAY)

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = date_p.add_run("2026-06-15")
set_run_fmt(r3, 12, color=DGRAY)

for _ in range(2):
    doc.add_paragraph()

spec_p = doc.add_paragraph()
spec_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r4 = spec_p.add_run("平台：PYNQ-Z2（Xilinx xc7z020clg400-1）\n"
                     "工具：Vitis HLS 2025.2 + Vivado 2025.2\n")
set_run_fmt(r4, 12)

status_p = doc.add_paragraph()
status_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r5 = status_p.add_run("狀態：csim PASS  |  cosim PASS  |  P&R Timing Met")
set_run_fmt(r5, 12, bold=True, color=GREEN)

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §1 背景與動機
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "1. 背景與動機：為什麼要這樣做？")

h2(doc, "1.1 什麼是高頻交易（HFT）？")
para(doc, "高頻交易（High-Frequency Trading, HFT）是一種在極短時間內（微秒到毫秒級別）自動下單買賣金融商品的策略。傳統交易中，人類觀察行情、分析訊號、決定下單，整個過程需要幾秒甚至幾分鐘。HFT 讓機器在人類反應之前就完成這整個流程。")
para(doc, "速度為何如此關鍵？金融市場是競爭性環境：若比對手晚 1 微秒接收到行情變動，你的訂單到達交易所前，好的價格已被搶走。在某些策略中，延遲降低 10 µs 就可以直接影響獲利能力。")

h2(doc, "1.2 為什麼使用 FPGA？")
para_mixed(doc, [
    ("FPGA（現場可程式化閘陣列）的三大核心優勢：", {"bold": True}),
])
bullet(doc, "確定性延遲（Deterministic Latency）：不受作業系統、記憶體垃圾回收、快取命中率影響，每次執行時間完全相同")
bullet(doc, "無軟體層（Bare Metal）：沒有 Linux、沒有 Python、沒有 PyTorch，訊號直接進晶片算完即出")
bullet(doc, "可客製化管線（Custom Pipeline）：根據模型架構設計專屬計算電路，不像 CPU 需要共用通用指令集")
para(doc, "GPU 雖然有高度並行能力，但其 batch 特性讓單次推論的延遲反而更高，更不適合 HFT 的逐筆推論需求。CPU 的 p99 延遲（779 µs）更會因為 OS 排程而產生尾端爆衝。")

h2(doc, "1.3 本專案目標")
para(doc, "在 xc7z020 學術 FPGA 上實作 PHC-lite 神經網路的推論核心，完成三層驗證：C simulation（csim）→ RTL co-simulation（cosim）→ Place & Route（P&R timing met）。")

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §2 模型架構
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "2. PHC-lite 模型架構")

h2(doc, "2.1 模型定義")
para(doc, "PHC-lite 是 PHCSSM（Predictive Hierarchical Continuous State-Space Model）的 16 維輕量版，為 FPGA 可行性展示而設計。它是一個帶記憶的神經網路（recurrent）：每步輸出的隱藏狀態（state_next[16]）會傳入下一步，讓模型能融合歷史走勢。")

para(doc, "輸入 / 輸出：")
bullet(doc, "輸入：x_t[16]（當前市場特徵）+ state_prev[16]（上一步隱藏狀態）")
bullet(doc, "輸出：readout[2]（買/賣方向訊號）+ state_next[16]（更新後隱藏狀態）")
bullet(doc, "資料型別：ap_fixed<24,12>（12 bit 整數 + 12 bit 小數）")

h2(doc, "2.2 計算流程（9 步驟）")
make_table(doc,
    ["步驟", "計算", "維度", "說明"],
    [
        ["Step 1", "ENC_LINEAR→LayerNorm→SiLU", "16→16→16→16", "投影輸入，正規化，激活"],
        ["Step 2", "GATE_LINEAR→Sigmoid",        "16→16→16",    "遺忘閘：控制保留多少舊記憶"],
        ["Step 3", "CANDIDATE_LINEAR→Tanh",      "16→16→16",    "候選記憶：這步要寫入的新資訊"],
        ["Step 4", "DECAY⊙state + gate⊙candidate","16→16",      "更新記憶：衰減舊狀態+加入新資訊"],
        ["Step 5", "Gain 計算→Sigmoid",           "16→16",       "增益：決定訊息傳遞強度"],
        ["Step 6", "SYNAPSE_LINEAR→SiLU",         "16→16→16",    "突觸訊息：模仿神經元訊號傳遞"],
        ["Step 7", "state + α⊙gain⊙message",      "16→16",       "自迴路更新：整合突觸訊息"],
        ["Step 8", "LayerNorm(loop_input)",        "16→16",       "最終正規化"],
        ["Step 9", "READOUT_LINEAR",               "16→2",        "輸出預測訊號"]
    ],
    [1.5, 4.0, 2.5, 5.5])

h2(doc, "2.3 為什麼是 Recurrent？")
para_mixed(doc, [
    ("重要限制：", {"bold": True}),
    ("t+1 步的輸入依賴 t 步的 state_next，存在資料相依性，不可跨步 pipeline。這是 recurrent 架構的根本約束，不是實作缺陷。每次計算必須等前一步完成才能開始。", {}),
])

h2(doc, "2.4 定點數型別選擇")
make_table(doc,
    ["型別", "定義", "整/小數位", "範圍", "用途"],
    [
        ["phc_fixed_t", "ap_fixed<24,12>", "12/12", "±2048，精度 1/4096", "I/O、權重"],
        ["phc_acc_t",   "ap_fixed<40,16>", "16/24", "±32768，高精度",      "乘法累加（防溢位）"],
        ["phc_var_t",   "ap_fixed<48,24>", "24/24", "±8M",               "LayerNorm 變異數"]
    ],
    [2.5, 3.5, 2.5, 4.0, 3.0])

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §3 工具鏈與設計策略
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "3. 工具鏈選擇與設計策略")

h2(doc, "3.1 HLS（High-Level Synthesis）")
para(doc, "傳統 FPGA 設計需要手寫 Verilog/VHDL（電路描述語言）。HLS 讓工程師用 C++ 描述計算邏輯，工具自動翻譯成硬體電路（RTL）。類比：HLS 對 FPGA 的關係，就像編譯器對 CPU——你寫高階語言，工具幫你轉成機器語言或電路。")

h2(doc, "3.2 手動 HLS vs hls4ml")
para(doc, "hls4ml 可以直接把 PyTorch 模型自動轉成 HLS，但預設策略是盡量展開（unroll）所有迴圈追求最低延遲，代價是資源用量極高。對 xc7z020（220 DSP）而言，hls4ml 的預設很可能讓 DSP 超過 100%，無法 P&R。手動 HLS 讓每個 pragma 可精確控制，在延遲和資源之間做出明確 tradeoff。")

h2(doc, "3.3 Tanh / Sigmoid 近似")
para(doc, "不用浮點的原因：HLS 看到標準 tanh() 就呼叫 Xilinx 雙精度浮點 IP 核心，單一核心需要 2000+ DSP，遠超過整個晶片上限（220 個）。")
para(doc, "解決方案：預計算 65 點定點 LUT（TANH_LUT[65]），範圍 [-4.0, +4.0]，步長 0.125，線性插值。Sigmoid 從 tanh 推導共用 LUT；SiLU = x × sigmoid(x) 同樣共用，不增加額外硬體。加上 ARRAY_PARTITION complete pragma 讓 LUT 存進 register，每個週期可同時讀取任意索引。")

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §4 優化歷程
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "4. 優化歷程：8 輪 HLS Pragma 調整")

h2(doc, "4.0 FPGA 資源背景")
make_table(doc,
    ["資源", "全名", "功能", "可用量"],
    [
        ["LUT",     "Look-Up Table",      "組合邏輯（AND/OR/加法）",     "53,200"],
        ["FF",      "Flip-Flop",          "儲存 1 bit 的暫存器",         "106,400"],
        ["DSP48E1", "Digital Signal Proc","高速乘法-累加（MAC）",         "220 ← 瓶頸"],
        ["BRAM_18K","Block RAM",          "18Kb 雙埠記憶體",             "280"]
    ],
    [2.5, 3.5, 5.0, 2.5])

h2(doc, "4.1 v1：初始版本（DSP 2636%）")
para_mixed(doc, [("根因：", {"bold": True}), ("標準 C tanh() 呼叫 Xilinx generic_tanh_double_s IP 核心，單一核心 2048 DSP，加上 sqrt/exp 核心，合計 > 5800 DSP（可用 220）。", {})])
para_mixed(doc, [("修復：", {"bold": True}), ("換用 65 點定點 LUT，完全消除浮點 IP 核心。", {})])

h2(doc, "4.2 v2～v5：嘗試限制展開（全部無效）")
para(doc, "v2 後 DSP 降至 1603%，主因是 #pragma HLS PIPELINE 強制展開所有 inner loop（WARNING [HLS 214-275]）。")
para(doc, "試過：移除 inner UNROLL（v3）、改 cyclic partition（v4）、改 II=297（v5）——全部無效。因為任何形式的 PIPELINE pragma 都會觸發這個行為，不是 II 值的問題，是 pragma 存在本身觸發的。")

h2(doc, "4.3 v6：刪除 PIPELINE pragma（DSP 247%）")
para(doc, "刪除所有 #pragma HLS PIPELINE，DSP 從 1603% 降到 247%。方向正確但仍超標。")

h2(doc, "4.4 v7：pipeline_loops 三連鎖（DSP 121%）")
para(doc, "移除 ARRAY_PARTITION 後，HLS 的 pipeline_loops 全域設定發動三連鎖：")
bullet(doc, "[HLS 214-376]：pipeline_loops 自動偵測 16-iter loop，自動加 PIPELINE")
bullet(doc, "[HLS 214-291]：自動 PIPELINE 強制 inner loop 完全展開（complete unroll）")
bullet(doc, "[HLS 214-449]：展開後索引成常數，自動 partition weight 陣列")
para(doc, "移除手動 pragma 無效，因工具會主動補回。需要明確的「禁止」指令。")

h2(doc, "4.5 v8：#pragma HLS PIPELINE off（最終，DSP 21%）")
para_mixed(doc, [("#pragma HLS PIPELINE off", {"bold": True}), (" 明確告知 HLS 不對此 loop pipeline，優先級高於 pipeline_loops 全域設定。三連鎖全部斷開。套用於所有 linear、layernorm、activation 函式的所有 loop。", {})])

h2(doc, "4.6 優化歷程總覽")
make_table(doc,
    ["版本", "關鍵動作", "DSP %", "LUT %", "FF %", "csim"],
    [
        ["v1",    "初始（double tanh）",          "2636%", "2073%", "441%", "未測"],
        ["v2",    "65點LUT tanh",                 "1603%", "839%",  "441%", "PASS"],
        ["v3~v5", "調整UNROLL/II（全部無效）",     "1603%", "839%",  "441%", "—"],
        ["v6",    "刪除 PIPELINE pragma",         "247%",  "113%",  "86%",  "PASS"],
        ["v7",    "移除 ARRAY_PARTITION",         "121%",  "77%",   "52%",  "—"],
        ["v8 ✓",  "#pragma HLS PIPELINE off",    "21%",   "47%",   "14%",  "PASS"]
    ],
    [1.8, 4.5, 1.5, 1.5, 1.5, 1.7])

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §5 驗證結果
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "5. 驗證結果")

h2(doc, "5.1 三層驗證鏈")
make_table(doc,
    ["層次", "工具", "驗證的是什麼", "通過代表什麼"],
    [
        ["C Simulation (csim)",       "Vitis HLS (clang)", "純 C++ 邏輯，無時序",        "演算法邏輯正確"],
        ["RTL Co-Simulation (cosim)", "ModelSim/Xsim",     "時脈精確的 RTL 閘級電路",    "RTL 與 C model bit-exact 一致"],
        ["Place & Route (P&R)",       "Vivado Implement",   "真實佈局在矽晶片的電路",    "資源在預算內、100MHz 時序收斂"]
    ],
    [4.0, 3.5, 4.5, 4.5])

h2(doc, "5.2 csim / cosim 結果")
make_table(doc,
    ["指標", "csim", "cosim", "說明"],
    [
        ["測試步數",    "25,600", "25,600",  "200 seq × 128 timestep"],
        ["Mismatches",  "0",      "0",       "全部通過"],
        ["max_abs_diff","0.00725275","0.00725275","完全相同（bit-exact）"],
        ["mean_abs_diff","0.000373", "0.000373",  "平均誤差極小"],
        ["Tolerance",   "1.0",    "1.0",     "允許最大誤差上限"]
    ],
    [3.5, 3.0, 3.0, 7.0])
para_mixed(doc, [
    ("csim 與 cosim 的 max_abs_diff 完全相同（0.00725275）：", {"bold": True}),
    ("代表 RTL 電路與 C model bit-level 一致，HLS 合成過程沒有引入任何額外捨入誤差。這是使用定點數的優勢——精度完全由位元寬決定。", {})
])

h2(doc, "5.3 Sign Agreement Rate（方向一致性）")
para_mixed(doc, [
    ("符號規則：", {"bold": True}),
    ("負值 → −1；零/正值 → +1。分析工具：03_sign_agreement.py，輸入：phc_step_cosim_out.csv", {})
])
make_table(doc,
    ["維度", "總步數", "同號", "異號", "Sign Agreement"],
    [
        ["readout[0]",   "25,600", "25,595", "5",  "99.9805%"],
        ["readout[1]",   "25,600", "25,589", "11", "99.9570%"],
        ["兩維聯合",      "25,600", "25,584", "16", "99.9375%"]
    ],
    [3.5, 2.5, 2.5, 2.5, 4.5])
para_mixed(doc, [
    ("關鍵解讀：", {"bold": True}),
    ("全部 16 個異號案例的 |ref| < 0.00054，代表模型在這些時步的輸出本身趨近零（模型無方向性意見）。若策略設有 confidence threshold，這些時步本身就會被過濾，不影響實際交易決策。", {})
])

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §6 硬體結果
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "6. 硬體結果（P&R 最終數字）")

h2(doc, "6.1 HLS Csynth 估算（v8）")
make_table(doc,
    ["時脈", "目標", "HLS 估算", "Latency"],
    [["ap_clk @ 100MHz", "10.00 ns", "7.147 ns", "8316 cycles = 83.16 µs"]],
    [3.0, 3.0, 3.0, 7.5])

h2(doc, "6.2 Post-Implementation 時序（Timing）")
make_table(doc,
    ["階段", "目標週期", "最差路徑", "WNS", "狀態"],
    [
        ["Post-Synthesis",       "10 ns", "9.729 ns", "+0.271 ns", "PASS"],
        ["Post-Implementation",  "10 ns", "9.980 ns", "+0.020 ns", "PASS ✓"]
    ],
    [4.0, 2.5, 2.5, 2.5, 2.0])
para_mixed(doc, [
    ("WNS = +0.020 ns：", {"bold": True}),
    ("timing met，但 margin 極薄。使用 Slow Corner（最保守製程/電壓/溫度條件）。22,941 個 endpoint 全部通過，TNS = 0.000 ns。", {}),
])
para_mixed(doc, [
    ("生產部署建議：", {"bold": True}),
    ("降頻至 90MHz → WNS 估計 +1.13 ns（符合工業 ≥ 1 ns 建議），Latency 從 83.16 µs 增至 92.4 µs（+11%）。", {})
])

h2(doc, "6.3 資源利用率（Utilization Routed）")
make_table(doc,
    ["資源", "csynth 估算", "P&R 實際", "可用量", "實際 %"],
    [
        ["Slice LUT",  "24,836", "7,555", "53,200",  "14.20%"],
        ["Slice FF",   "14,518", "8,217", "106,400",  "7.72%"],
        ["DSP48E1",    "46",     "47",    "220",      "21.36%"],
        ["BRAM_18K",   "6",      "6",     "280",       "2.14%"]
    ],
    [3.0, 2.5, 2.5, 2.5, 2.0])
para_mixed(doc, [
    ("LUT 從 24,836 降至 7,555（降幅 70%）：", {"bold": True}),
    ("Vivado synthesis 執行邏輯合併、常數傳播、死碼消除。序列執行（PIPELINE off）讓大量邏輯可時間複用，Vivado 優化空間大。此降幅完全符合預期，不是數字矛盾。", {})
])

h2(doc, "6.4 功耗分析")
make_table(doc,
    ["項目", "數值", "說明"],
    [
        ["Total On-Chip Power", "0.357 W",   "Dynamic + Static"],
        ["Dynamic Power",       "0.251 W",   "電路切換消耗"],
        ["Device Static Power", "0.107 W",   "漏電流"],
        ["Junction Temperature","29.1 °C",   "@ 25°C 環境溫度"],
        ["Max Ambient Temp",    "80.9 °C",   "熱安全上限"],
        ["Confidence Level",    "Medium",    "內部節點活動率使用預設值（0.5）"]
    ],
    [4.0, 3.0, 9.5])

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §7 效能比較
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "7. 效能比較：FPGA vs CPU")

h2(doc, "7.1 CPU Baseline 設定")
make_table(doc,
    ["項目", "設定"],
    [
        ["平台",   "Windows 11，Intel CPU"],
        ["模型",   "PHC-lite，PyTorch float32（phc_stp_step_hls_seed2.pt）"],
        ["計時",   "time.perf_counter()，warm-up 100次，計時 10,000 次"],
        ["模式",   "torch.no_grad()（關閉梯度計算），batch_size=1"],
        ["NEW5",   "new5_used=False（不使用保留測試集）"]
    ],
    [3.0, 13.5])

h2(doc, "7.2 延遲比較")
make_table(doc,
    ["指標", "FPGA（v8）", "CPU（no_grad）", "說明"],
    [
        ["Latency（固定）",  "83.16 µs", "—",          "FPGA 完全確定性，每次相同"],
        ["Latency（mean）",  "83.16 µs", "117.5 µs",   "FPGA 快 1.4×"],
        ["Latency（median）","83.16 µs", "84.2 µs",    "幾乎相同"],
        ["Latency（min）",   "83.16 µs", "61.2 µs",    "CPU min 較低（JIT cache hit）"],
        ["Latency（p99）",   "83.16 µs", "778.8 µs",   "FPGA 優勢 9.4×"],
        ["Jitter",           "0",        "高（OS/GC）", "FPGA 核心優勢"],
        ["Determinism",      "完全確定",  "不確定",     "HFT 最重要特性"]
    ],
    [3.5, 2.5, 2.5, 8.0])

h2(doc, "7.3 正確的比較視角")
para_mixed(doc, [
    ("FPGA 的優勢在 tail latency 和確定性，不在 mean latency。", {"bold": True})
])
para(doc, "CPU p99 = 779 µs 代表每 100 次執行就有一次超過 779 µs，那一次等同於錯失交易機會。FPGA 不論何時執行都是 83.16 µs。CPU min = 61 µs 是 JIT cache 命中的最佳情況，在真實系統中幾乎不可能持續出現。")
para_mixed(doc, [
    ("誠實聲明：", {"bold": True}),
    ("83 µs 不滿足 HFT L1 < 10 µs 需求。本專案展示 FPGA 推論可行性與 HLS 設計方法論，不聲稱改善 PnL 或達到生產水準。", {})
])

page_break(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# §8 結論與改進路徑
# ═══════════════════════════════════════════════════════════════════════════════
h1(doc, "8. 結論與改進路徑")

h2(doc, "8.1 改進路徑（初步估算，未經 synthesis 驗證）")
make_table(doc,
    ["方法", "預期 Latency", "代價"],
    [
        ["INT16 量化（ap_fixed<16,8>）",   "~25 µs",  "精度降低，需重新驗證 sign agreement"],
        ["Hidden dim 縮至 8",               "~8 µs",   "模型表達能力降低"],
        ["換用 Alveo U50（5,952 DSP）",     "< 5 µs",  "成本大幅上升，需重新 P&R"]
    ],
    [5.5, 2.5, 8.5])

h2(doc, "8.2 結論")
bullet(doc, "8 輪系統性 HLS pragma 優化，DSP 從 2636% → 21%（可用量的五分之一）")
bullet(doc, "確認 HLS PIPELINE pragma 觸發 cascade unroll 根本機制，以及 pipeline_loops 三連鎖")
bullet(doc, "csim + cosim 25,600 步全部通過，max_abs_diff = 0.00725，sign agreement 99.9375%")
bullet(doc, "P&R timing met（WNS +0.020 ns，100 MHz），所有資源在預算內")
bullet(doc, "功耗 0.357 W，Junction Temp 29.1°C，熱餘量充足")
bullet(doc, "與 CPU 相比，FPGA 實現 jitter-free 確定性延遲，p99 改善 9.4 倍")

# ─── 儲存 ─────────────────────────────────────────────────────────────────────
output_path = "PHC_lite_HLS_Report.docx"
doc.save(output_path)
print(f"✓ {output_path} 生成成功")
