"""Build the professor-facing presentation deck from the latest artifacts."""

from __future__ import annotations

import argparse
from datetime import date, timedelta
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
MEETING_DIR = ROOT / "meeting" / "w5"
CHART_DIR = MEETING_DIR / "w5_chart"
OUTPUT_DIR = MEETING_DIR

FONT_SANS = "Microsoft JhengHei"
FONT_MONO = "Consolas"

BG = RGBColor(246, 243, 238)
NAVY = RGBColor(15, 33, 48)
INK = RGBColor(31, 40, 51)
MUTED = RGBColor(101, 115, 126)
TEAL = RGBColor(0, 155, 149)
TEAL_DARK = RGBColor(0, 112, 110)
AMBER = RGBColor(226, 152, 52)
RED = RGBColor(184, 57, 49)
GREEN = RGBColor(71, 160, 84)
LIGHT = RGBColor(226, 234, 235)
SOFT_RED = RGBColor(250, 231, 228)
SOFT_GREEN = RGBColor(229, 244, 233)
SOFT_TEAL = RGBColor(229, 245, 244)
SOFT_AMBER = RGBColor(251, 241, 222)


def _set_run_font(run, font_name: str, size: float, color: RGBColor, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    r_pr.set(qn("a:lang"), "zh-TW")
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set("typeface", font_name)


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, fill_color, line_color=None, radius=True, transparency=0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = fill_color
    fill.transparency = transparency
    line = shape.line
    if line_color is None:
        line.fill.background()
    else:
        line.color.rgb = line_color
        line.width = Pt(1.25)
    return shape


def add_line(slide, x1, y1, x2, y2, color: RGBColor, width: float = 2.0):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    paragraphs,
    font_name=FONT_SANS,
    font_size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    margin=0.12,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign

    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]

    for idx, entry in enumerate(paragraphs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        text = entry["text"] if isinstance(entry, dict) else str(entry)
        p.text = text
        p.alignment = entry.get("align", align) if isinstance(entry, dict) else align
        p.level = entry.get("level", 0) if isinstance(entry, dict) else 0
        if isinstance(entry, dict):
            p.space_after = Pt(entry.get("space_after", 0))
            p.space_before = Pt(entry.get("space_before", 0))
            p.line_spacing = entry.get("line_spacing", 1.05)
        run = p.runs[0]
        _set_run_font(
            run,
            entry.get("font_name", font_name) if isinstance(entry, dict) else font_name,
            entry.get("font_size", font_size) if isinstance(entry, dict) else font_size,
            entry.get("color", color) if isinstance(entry, dict) else color,
            entry.get("bold", bold) if isinstance(entry, dict) else bold,
        )
    return box


def add_title(slide, title: str, subtitle: str | None = None, dark=False):
    title_color = RGBColor(255, 255, 255) if dark else NAVY
    sub_color = RGBColor(226, 234, 235) if dark else MUTED
    add_textbox(
        slide,
        Inches(0.6),
        Inches(0.45),
        Inches(8.8),
        Inches(0.8),
        [{"text": title, "font_size": 24, "bold": True, "color": title_color}],
    )
    if subtitle:
        add_textbox(
            slide,
            Inches(0.62),
            Inches(1.08),
            Inches(7.8),
            Inches(0.45),
            [{"text": subtitle, "font_size": 10.5, "color": sub_color}],
        )


def add_chip(slide, x, y, text, fill_color, text_color=RGBColor(255, 255, 255), w=None):
    width = w or Inches(1.6)
    shape = add_rect(slide, x, y, width, Inches(0.42), fill_color, None, radius=True)
    add_textbox(
        slide,
        x,
        y + Inches(0.01),
        width,
        Inches(0.34),
        [{"text": text, "font_size": 10.5, "color": text_color, "bold": True, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0.02,
    )
    return shape


def add_image_fit(slide, path: Path, x, y, w, h):
    with Image.open(path) as img:
        img_w, img_h = img.size
    box_ratio = w / h
    img_ratio = img_w / img_h
    if img_ratio > box_ratio:
        final_w = w
        final_h = w / img_ratio
        top = y + (h - final_h) / 2
        left = x
    else:
        final_h = h
        final_w = h * img_ratio
        left = x + (w - final_w) / 2
        top = y
    slide.shapes.add_picture(str(path), left, top, width=int(final_w), height=int(final_h))


def add_footer(slide, text: str, dark=False):
    color = RGBColor(220, 228, 231) if dark else MUTED
    add_textbox(
        slide,
        Inches(0.62),
        Inches(7.04),
        Inches(12.0),
        Inches(0.22),
        [{"text": text, "font_size": 8.5, "color": color}],
    )


def add_metric_card(slide, x, y, w, title, value, fill_color):
    add_rect(slide, x, y, w, Inches(0.92), fill_color, None)
    add_textbox(
        slide,
        x + Inches(0.12),
        y + Inches(0.1),
        w - Inches(0.24),
        Inches(0.2),
        [{"text": title, "font_size": 10, "color": MUTED, "bold": True}],
        margin=0.0,
    )
    add_textbox(
        slide,
        x + Inches(0.12),
        y + Inches(0.33),
        w - Inches(0.24),
        Inches(0.38),
        [{"text": value, "font_size": 19, "color": NAVY, "bold": True}],
        margin=0.0,
    )


def build_presentation(presentation_date: date) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, NAVY)
    add_rect(slide, Inches(-0.5), Inches(5.45), Inches(5.5), Inches(3.5), TEAL_DARK, None, transparency=0.15)
    add_rect(slide, Inches(9.4), Inches(-0.7), Inches(4.5), Inches(3.6), AMBER, None, transparency=0.2)
    add_textbox(
        slide,
        Inches(0.65),
        Inches(0.55),
        Inches(2.5),
        Inches(0.32),
        [{"text": f"TOMORROW BRIEFING  |  {presentation_date.isoformat()}", "font_size": 9.5, "color": RGBColor(216, 224, 227), "bold": True}],
        margin=0.0,
    )
    add_textbox(
        slide,
        Inches(0.65),
        Inches(1.2),
        Inches(8.5),
        Inches(1.3),
        [
            {"text": "HFT Project 指導教授進度簡報", "font_size": 28, "bold": True, "color": RGBColor(255, 255, 255)},
            {"text": "Week 4-5 策略轉向、FPGA 價值證明與 Week 6 執行計畫", "font_size": 15, "color": RGBColor(224, 232, 235), "space_before": 6},
        ],
    )
    add_textbox(
        slide,
        Inches(0.68),
        Inches(2.75),
        Inches(7.6),
        Inches(1.2),
        [
            {
                "text": "主線訊息：我們不是臨時改方向，而是把專案從「預測準確率競賽」重新拉回最能量化展示 FPGA 低延遲價值的策略路徑。",
                "font_size": 17,
                "color": RGBColor(240, 245, 246),
            }
        ],
    )
    add_chip(slide, Inches(0.68), Inches(5.4), "Pipeline Ready", TEAL, w=Inches(1.65))
    add_chip(slide, Inches(2.5), Inches(5.4), "Model Frozen", GREEN, w=Inches(1.55))
    add_chip(slide, Inches(4.2), Inches(5.4), "Latency Proof", AMBER, text_color=NAVY, w=Inches(1.65))
    add_rect(slide, Inches(8.6), Inches(1.2), Inches(3.95), Inches(4.85), RGBColor(19, 45, 63), RGBColor(78, 112, 132))
    add_textbox(
        slide,
        Inches(8.92),
        Inches(1.55),
        Inches(3.2),
        Inches(0.35),
        [{"text": "明天簡報的三件事", "font_size": 14, "color": RGBColor(255, 255, 255), "bold": True}],
    )
    for idx, item in enumerate(
        [
            "1. Week 4-5 已完成端到端資料、模型與事件回測管線。",
            "2. 證明真正有價值的不是更保守的預測，而是更快的撤單。",
            "3. Week 6 進入 CME 擴充、C++ 推論與 HLS handoff 準備。",
        ]
    ):
        add_textbox(
            slide,
            Inches(8.95),
            Inches(2.02 + idx * 1.02),
            Inches(3.08),
            Inches(0.74),
            [{"text": item, "font_size": 13.5, "color": RGBColor(226, 234, 237)}],
        )
    add_footer(slide, "Sources: meeting/w5_meeting_outline.md, schedule.md, guide.md", dark=True)

    # Slide 2
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, BG)
    add_title(slide, "Workflow 對齊：這次 pivot 是把主線拉回「可證明 FPGA 價值」", "對照 schedule.md 的雙軌 workflow，明天簡報聚焦 Hybrid MM / latency-sensitive 這一支")

    phase_x = [0.7, 3.02, 5.34, 7.66, 10.0]
    phases = [
        ("W1-4", "特徵工程 + TCN baseline", SOFT_TEAL),
        ("W5", "Pivot to Hybrid MM", SOFT_AMBER),
        ("W6", "校準與模型選型", SOFT_TEAL),
        ("W7-12", "INT8 / C++ / HLS", SOFT_GREEN),
        ("W13-16", "Simulator + FPGA value report", LIGHT),
    ]
    for idx, (tag, label, fill_color) in enumerate(phases):
        add_rect(slide, Inches(phase_x[idx]), Inches(1.45), Inches(1.95), Inches(0.88), fill_color, RGBColor(206, 214, 217))
        add_textbox(
            slide,
            Inches(phase_x[idx] + 0.12),
            Inches(1.58),
            Inches(1.72),
            Inches(0.52),
            [
                {"text": tag, "font_size": 10, "color": MUTED, "bold": True},
                {"text": label, "font_size": 12.5, "color": NAVY, "bold": True, "space_before": 2},
            ],
            margin=0.0,
        )
        if idx < len(phases) - 1:
            add_line(slide, Inches(phase_x[idx] + 1.95), Inches(1.89), Inches(phase_x[idx + 1]), Inches(1.89), TEAL_DARK, 2.6)

    add_rect(slide, Inches(0.72), Inches(2.75), Inches(5.95), Inches(2.5), RGBColor(251, 252, 252), RGBColor(221, 228, 231))
    add_textbox(
        slide,
        Inches(0.98),
        Inches(3.0),
        Inches(5.35),
        Inches(0.35),
        [{"text": "Directional microstructure alpha", "font_size": 16, "color": NAVY, "bold": True}],
    )
    add_chip(slide, Inches(4.85), Inches(2.96), "Reference Track", TEAL_DARK, w=Inches(1.45))
    add_textbox(
        slide,
        Inches(0.98),
        Inches(3.42),
        Inches(5.35),
        Inches(1.5),
        [
            {"text": "目標：優化交易品質，讓現有 direction-aware baseline 靠近穩定正 PnL。", "font_size": 13},
            {"text": "現況：schedule 中最接近 trading viability 的結果為 val-calibrated test net PnL = -3 ticks。", "font_size": 13, "space_before": 6},
            {"text": "角色：保留為 reference baseline，後續用來比較共享 TCN 是否值得。", "font_size": 13, "space_before": 6},
        ],
    )

    add_rect(slide, Inches(6.95), Inches(2.75), Inches(5.65), Inches(2.5), RGBColor(252, 251, 248), RGBColor(221, 228, 231))
    add_textbox(
        slide,
        Inches(7.2),
        Inches(3.0),
        Inches(5.05),
        Inches(0.35),
        [{"text": "Hybrid MM / queue-reaction", "font_size": 16, "color": NAVY, "bold": True}],
    )
    add_chip(slide, Inches(11.05), Inches(2.96), "Meeting Focus", AMBER, text_color=NAVY, w=Inches(1.35))
    add_textbox(
        slide,
        Inches(7.2),
        Inches(3.42),
        Inches(5.05),
        Inches(1.5),
        [
            {"text": "目標：把 TCN 用在 latency-sensitive 的 quote / cancel 決策。", "font_size": 13},
            {"text": "現況：已形成 killer charts，可以直接展示『低延遲撤單』對存活的決定性影響。", "font_size": 13, "space_before": 6},
            {"text": "角色：最能對齊 FPGA 展示價值，也是明天簡報的核心故事線。", "font_size": 13, "space_before": 6},
        ],
    )

    add_rect(slide, Inches(0.72), Inches(5.65), Inches(11.9), Inches(0.92), SOFT_RED, None)
    add_textbox(
        slide,
        Inches(0.95),
        Inches(5.9),
        Inches(11.4),
        Inches(0.36),
        [{"text": "結論：明天不是在報告「又換一個想法」，而是在報告我們已經找到最能把 FPGA 低延遲優勢量化出來的正確展示路徑。", "font_size": 15, "color": RED, "bold": True}],
    )
    add_footer(slide, "Workflow reference: schedule.md (Updated Schedule, Week 4.11 to Week 16)")

    # Slide 3
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, BG)
    add_title(slide, "Week 4-5 里程碑：端到端管線已經完成，TCN 也已具備硬體移交條件", "meeting/w5_meeting_outline.md 投影片 1 + guide.md 整體規劃")
    add_metric_card(slide, Inches(0.72), Inches(1.42), Inches(1.8), "INT8 權重", "4.8 KB", SOFT_TEAL)
    add_metric_card(slide, Inches(2.7), Inches(1.42), Inches(1.8), "Activation", "12.5 KB", SOFT_GREEN)
    add_metric_card(slide, Inches(4.68), Inches(1.42), Inches(1.8), "FP32 參考", "19.2 KB", SOFT_AMBER)

    add_rect(slide, Inches(0.72), Inches(2.58), Inches(5.85), Inches(3.65), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_textbox(
        slide,
        Inches(0.98),
        Inches(2.9),
        Inches(5.3),
        Inches(2.6),
        [
            {"text": "已完成進度", "font_size": 15, "color": NAVY, "bold": True},
            {"text": "• 資料處理管道、無因次特徵、TCN 建模與 event-driven backtester 已串成完整 workflow。", "font_size": 13.5, "space_before": 8},
            {"text": "• `frozen_tcn_best_model.pt` 已凍結，可作為硬體端 handoff 起點。", "font_name": FONT_MONO, "font_size": 12.8, "space_before": 6},
            {"text": "• 模型記憶體占用與 MAC 規模都夠小，硬體可行性已被量化證明。", "font_size": 13.5, "space_before": 6},
            {"text": "• 下一個真正的瓶頸不在模型大小，而在『哪一種策略最吃低延遲』。", "font_size": 13.5, "space_before": 6, "bold": True, "color": TEAL_DARK},
        ],
    )
    add_rect(slide, Inches(6.82), Inches(1.42), Inches(5.73), Inches(4.95), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_image_fit(slide, CHART_DIR / "p5_03_hardware_footprint.png", Inches(7.05), Inches(1.72), Inches(5.25), Inches(4.25))
    add_footer(slide, "Chart source: meeting/w5/w5_chart/p5_03_hardware_footprint.png")

    # Slide 4
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, BG)
    add_title(slide, "死胡同：過度防禦的 Quote-off，把問題藏起來而不是解掉", "meeting/w5_meeting_outline.md 投影片 2")
    add_rect(slide, Inches(0.72), Inches(1.38), Inches(7.15), Inches(5.35), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_image_fit(slide, CHART_DIR / "p5_01b_latency_failed_initial.png", Inches(0.9), Inches(1.62), Inches(6.8), Inches(4.9))

    add_rect(slide, Inches(8.1), Inches(1.38), Inches(4.48), Inches(5.35), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_textbox(
        slide,
        Inches(8.35),
        Inches(1.72),
        Inches(4.0),
        Inches(3.45),
        [
            {"text": "觀察重點", "font_size": 15, "color": NAVY, "bold": True},
            {"text": "• `delay=0` 時虧損 -9.0 ticks，但只做了 2 筆交易。", "font_name": FONT_MONO, "font_size": 12.7, "space_before": 8},
            {"text": "• `delay=4` 時看似不虧，只是因為 0 trades，根本沒有真的參與市場。", "font_name": FONT_MONO, "font_size": 12.7, "space_before": 6},
            {"text": "• 問題不是 TCN 預測能力不夠，而是早期把 Quote-off 疊得太厚，直接把策略變成『不出手』。", "font_size": 13.4, "space_before": 6},
            {"text": "• 這種設計無法體現硬體速度優勢，也無法回答 FPGA 值不值得。", "font_size": 13.4, "space_before": 6, "bold": True, "color": RED},
        ],
    )
    add_rect(slide, Inches(8.35), Inches(5.38), Inches(3.95), Inches(0.9), SOFT_RED, None)
    add_textbox(
        slide,
        Inches(8.52),
        Inches(5.61),
        Inches(3.6),
        Inches(0.38),
        [{"text": "如果機器人連單都不敢掛，FPGA 再快也沒有展示舞台。", "font_size": 12.8, "bold": True, "color": RED}],
    )
    add_footer(slide, "Chart source: meeting/w5/w5_chart/p5_01b_latency_failed_initial.png")

    # Slide 5
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, BG)
    add_title(slide, "第一性原理 Pivot：不再追求更神準的預言，而是讓撤單速度真正變值錢", "meeting/w5_meeting_outline.md 投影片 3")
    add_rect(slide, Inches(0.72), Inches(1.35), Inches(11.9), Inches(0.9), SOFT_AMBER, None)
    add_textbox(
        slide,
        Inches(0.95),
        Inches(1.62),
        Inches(11.4),
        Inches(0.34),
        [{"text": "Net PnL = Spread Capture - Adverse Selection", "font_name": FONT_MONO, "font_size": 19, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER}],
        font_name=FONT_MONO,
        align=PP_ALIGN.CENTER,
    )

    step_x = [0.95, 4.35, 7.75]
    step_text = [
        ("1. Dumb Quoting", "先恢復常態雙邊掛單，讓策略有機會穩定吃 spread。", SOFT_TEAL),
        ("2. TCN Toxicity Radar", "把 TCN 角色改成盤面毒性偵測器，而不是方向神諭。", SOFT_GREEN),
        ("3. FPGA Smart Canceling", "只要偵測到失衡，就把 latency 優勢用在極速撤單。", SOFT_AMBER),
    ]
    for idx, (headline, body, fill_color) in enumerate(step_text):
        add_rect(slide, Inches(step_x[idx]), Inches(2.65), Inches(2.9), Inches(2.1), fill_color, RGBColor(217, 224, 226))
        add_textbox(
            slide,
            Inches(step_x[idx] + 0.16),
            Inches(2.88),
            Inches(2.58),
            Inches(1.52),
            [
                {"text": headline, "font_size": 15, "color": NAVY, "bold": True, "align": PP_ALIGN.CENTER},
                {"text": body, "font_size": 13, "color": INK, "space_before": 10, "align": PP_ALIGN.CENTER},
            ],
            align=PP_ALIGN.CENTER,
        )
        if idx < 2:
            add_line(slide, Inches(step_x[idx] + 2.9), Inches(3.7), Inches(step_x[idx + 1]), Inches(3.7), TEAL_DARK, 2.5)

    add_rect(slide, Inches(0.95), Inches(5.28), Inches(11.36), Inches(1.0), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_textbox(
        slide,
        Inches(1.2),
        Inches(5.58),
        Inches(10.8),
        Inches(0.34),
        [{"text": "關鍵重構：重點不再是「猜方向多準」，而是「危險來時能不能比一般軟體更早把單撤掉」。", "font_size": 15, "bold": True, "color": TEAL_DARK, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Storyline source: meeting/w5_meeting_outline.md (First Principles Pivot)")

    # Slide 6
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, BG)
    add_title(slide, "Killer Charts：同一套策略，只改 latency，PnL 與 adverse loss 立刻翻盤", "meeting/w5_meeting_outline.md 投影片 4")
    add_rect(slide, Inches(0.72), Inches(1.36), Inches(5.9), Inches(4.2), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_rect(slide, Inches(6.73), Inches(1.36), Inches(5.9), Inches(4.2), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
    add_image_fit(slide, CHART_DIR / "p5_01_latency_expected_pnl.png", Inches(0.9), Inches(1.58), Inches(5.55), Inches(3.82))
    add_image_fit(slide, CHART_DIR / "p5_02_adverse_selection_loss.png", Inches(6.91), Inches(1.58), Inches(5.55), Inches(3.82))

    add_rect(slide, Inches(0.72), Inches(5.78), Inches(3.7), Inches(0.9), SOFT_GREEN, None)
    add_rect(slide, Inches(4.58), Inches(5.78), Inches(3.8), Inches(0.9), SOFT_RED, None)
    add_rect(slide, Inches(8.53), Inches(5.78), Inches(4.08), Inches(0.9), SOFT_AMBER, None)
    add_textbox(
        slide,
        Inches(0.9),
        Inches(6.02),
        Inches(3.3),
        Inches(0.34),
        [{"text": "PnL: +120.5 → +10.0 → -85.0 → -150.0 ticks", "font_name": FONT_MONO, "font_size": 12.5, "bold": True, "color": GREEN}],
        font_name=FONT_MONO,
    )
    add_textbox(
        slide,
        Inches(4.76),
        Inches(6.02),
        Inches(3.4),
        Inches(0.34),
        [{"text": "Loss: 5 → 45 → 110 → 180 ticks", "font_name": FONT_MONO, "font_size": 12.5, "bold": True, "color": RED}],
        font_name=FONT_MONO,
    )
    add_textbox(
        slide,
        Inches(8.72),
        Inches(5.93),
        Inches(3.7),
        Inches(0.45),
        [{"text": "解讀：不是 alpha 消失，而是撤單慢一步就被毒性流量吃掉。", "font_size": 12.8, "bold": True, "color": NAVY}],
    )
    add_footer(slide, "Chart sources: meeting/w5/w5_chart/p5_01_latency_expected_pnl.png, meeting/w5/w5_chart/p5_02_adverse_selection_loss.png")

    # Slide 7
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, BG)
    add_title(slide, "Week 6 執行計畫：把 killer story 往可交付的硬體與跨市場驗證推進", "對齊 schedule.md 的 Week 6 deliverable：model-selection memo + hardware-friendly path")
    columns = [
        (
            Inches(0.72),
            "Data / Market",
            [
                "撰寫 CME parser，先接 ES / NQ / CL 歷史資料。",
                "沿用無因次特徵工程，驗證跨市場泛化能力。",
                "保留台指期結果，建立跨市場可比較基準。",
            ],
            SOFT_TEAL,
        ),
        (
            Inches(4.47),
            "Software / Hardware",
            [
                "匯出 frozen TCN 權重，整理固定 feature ordering 與 tensor shape。",
                "開始建立 C++ inference skeleton，為後續 HLS 轉換鋪路。",
                "明確定義低延遲 cancel path 的 I/O 介面。",
            ],
            SOFT_GREEN,
        ),
        (
            Inches(8.22),
            "Selection / Deliverable",
            [
                "比較 trading quality、turnover、latency sensitivity、hardware friendliness。",
                "判斷要共用一顆 TCN，還是 directional / MM 分流兩顆模型。",
                "整理 Week 6 model-selection memo，作為下次會議交付。",
            ],
            SOFT_AMBER,
        ),
    ]
    for x, header, lines, fill_color in columns:
        add_rect(slide, x, Inches(1.55), Inches(3.18), Inches(4.65), RGBColor(252, 252, 252), RGBColor(221, 228, 231))
        add_rect(slide, x, Inches(1.55), Inches(3.18), Inches(0.62), fill_color, None)
        add_textbox(
            slide,
            x + Inches(0.16),
            Inches(1.73),
            Inches(2.9),
            Inches(0.24),
            [{"text": header, "font_size": 15, "bold": True, "color": NAVY}],
            margin=0.0,
        )
        y = 2.35
        for line in lines:
            add_textbox(
                slide,
                x + Inches(0.18),
                Inches(y),
                Inches(2.82),
                Inches(0.75),
                [{"text": f"• {line}", "font_size": 13.2, "color": INK}],
            )
            y += 1.0

    add_rect(slide, Inches(0.72), Inches(6.45), Inches(11.9), Inches(0.62), NAVY, None)
    add_textbox(
        slide,
        Inches(0.95),
        Inches(6.62),
        Inches(11.45),
        Inches(0.22),
        [{"text": "下次會議希望帶來：CME 資料接收器初版 + C++ TCN 推論骨架 + model-selection memo。", "font_size": 12.8, "color": RGBColor(255, 255, 255), "bold": True, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )
    add_footer(slide, "Execution plan references: schedule.md Week 6, guide.md, meeting/w5_meeting_outline.md 投影片 5", dark=True)

    props = prs.core_properties
    props.title = "HFT Project 指導教授進度簡報"
    props.subject = "Week 4-5 策略轉向與 Week 6 計畫"
    props.author = "Codex"
    props.keywords = "HFT, FPGA, TCN, Hybrid MM, Professor Meeting"
    props.comments = "Generated from meeting materials and workflow documents."

    return prs


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate professor meeting PPTX.")
    parser.add_argument("--date", dest="presentation_date", help="Presentation date in YYYY-MM-DD format.")
    parser.add_argument("--output", dest="output", help="Output PPTX path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.presentation_date:
        presentation_date = date.fromisoformat(args.presentation_date)
    else:
        presentation_date = date.today() + timedelta(days=1)

    output = Path(args.output) if args.output else OUTPUT_DIR / f"{presentation_date.isoformat()}_professor_update_week5_pivot.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = build_presentation(presentation_date)
    prs.save(output)
    print(output)


if __name__ == "__main__":
    main()
